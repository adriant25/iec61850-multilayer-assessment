"""
Build an isolated DES run folder: copy of Simulador code + Inputs taken from a CASO folder.
Original case files and Simulador/ are never modified.

usage: make_des_run.py <run_dir> <N> <pb_switch_file_suffix> [fix_bay_responses]
"""
import os, shutil, sys, glob
import openpyxl

# DES_CASES_DIR: folder with the CASO<N>BAHIA(S) input workbooks of each size;
# DES_SIMULATOR_DIR: simulator code (iec61850-network-simulator, tag v1.0-paper).
SIM = os.environ['DES_CASES_DIR']
SIMULATOR = os.environ['DES_SIMULATOR_DIR']
run_dir, n, sw_suffix = sys.argv[1], int(sys.argv[2]), sys.argv[3]
fix = len(sys.argv) > 4 and sys.argv[4] == 'fix'

case = glob.glob(os.path.join(SIM, f'CASO{n}BAHIA*'))[0]
p = lambda name: os.path.join(case, f'{n}BAHIA{name}.xlsx')

if os.path.exists(run_dir):
    shutil.rmtree(run_dir)
shutil.copytree(SIMULATOR, run_dir,
                ignore=shutil.ignore_patterns('Inputs', 'RespaldoInputs', 'Results',
                                              'txt_exports', '__pycache__', '.git'))
inp = os.path.join(run_dir, 'Inputs')
os.makedirs(inp)
mapping = {
    'goose_faults': p('FALLAs'), 'goose_pb_traffic': p('FLUJOSGOOSEPB'),
    'goose_responses': p('RAFAGAS'), 'goose_sb_traffic': p('FLUJOSGOOSESB'),
    'pb_switches': os.path.join(case, f'{n}BAHIASWPB{sw_suffix}.xlsx'),
    'pb_vlans': p('VLANPB'), 'sb_switches': p('SWSB'), 'sb_vlans': p('VLANSB'),
    'sv_traffic': p('FLUJOSSVPB'),
}
for dst, src in mapping.items():
    shutil.copy2(src, os.path.join(inp, dst + '.xlsx'))
    print(f'{dst:18s} <- {os.path.basename(src)}')

# Stubs for GUI-only modules (plots are declined at the prompts)
os.makedirs(os.path.join(run_dir, 'PyQt5'))
open(os.path.join(run_dir, 'PyQt5', '__init__.py'), 'w').close()
with open(os.path.join(run_dir, 'PyQt5', 'QtWidgets.py'), 'w') as f:
    f.write('class _Stub:\n    def __init__(self, *a, **k): pass\n'
            'def __getattr__(name):\n    return _Stub\n')
with open(os.path.join(run_dir, 'pyvista.py'), 'w') as f:
    f.write('# stub: 3D visualisation disabled for batch runs\n')

if fix:
    # The N=9 response table has the MU1B9/MU2B9 rows copied from bay 8.
    # Correct rows, derived from the flow table (same pattern as every other bay):
    #   MU1B9: trigger 128 (PP1B9L trip) -> activate 132;134;135
    #   MU2B9: trigger 129 (PP2B9L trip) -> activate 133;136;137
    path = os.path.join(inp, 'goose_responses.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(hdr)}
    for r in range(2, ws.max_row + 1):
        eq = ws.cell(r, ci['EQUIPMENT']).value
        if eq == 'MU1B9':
            ws.cell(r, ci['TRIGGERING FLOW ID']).value = 128
            ws.cell(r, ci['ACTIVATED FLOW IDS IN PB']).value = '132;134;135'
        elif eq == 'MU2B9':
            ws.cell(r, ci['TRIGGERING FLOW ID']).value = 129
            ws.cell(r, ci['ACTIVATED FLOW IDS IN PB']).value = '133;136;137'
    wb.save(path)
    print('goose_responses: MU1B9/MU2B9 rows corrected')

    # The N=9 SW1 VLAN sheet lacks the bay-9 VLANs used by GOOSE flows that
    # cross the trunk into SW1 (114-121), so SW1 floods those frames to every
    # port. Add each missing VLAN with trunk port + SW1 destination ports,
    # priority taken from the flow table.
    import pandas as pd, re
    sw = pd.read_excel(os.path.join(inp, 'pb_switches.xlsx'), sheet_name=None)
    port_of, trunk = {}, {}
    for s, df in sw.items():
        for port, eq in zip(df['PORT'], df['EQUIPMENT'].astype(str)):
            if 'TRUNK' in eq:
                trunk[s] = port
            else:
                port_of[eq.replace(' PB', '')] = (s, port)
    go = pd.read_excel(os.path.join(inp, 'goose_pb_traffic.xlsx'))
    vpath = os.path.join(inp, 'pb_vlans.xlsx')
    wb = openpyxl.load_workbook(vpath)
    for s in wb.sheetnames:
        ws = wb[s]
        have = {int(ws.cell(r, 1).value) for r in range(2, ws.max_row + 1)
                if ws.cell(r, 1).value is not None}
        need = {}
        for _, r in go.iterrows():
            vid = int(r['VLAN'])
            src_sw = port_of[str(r['SOURCE'])][0]
            dsts = [d.strip() for d in str(r['DESTINATION(S)']).split(';')]
            local = [port_of[d][1] for d in dsts if port_of[d][0] == s]
            if vid in have or not local or src_sw == s:
                continue
            ports = need.setdefault(vid, [{trunk[s]}, int(r['PRIORITY'])])[0]
            ports.update(local)
        for vid, (ports, prio) in sorted(need.items()):
            ws.append([vid, ';'.join(str(p) for p in sorted(ports)), prio])
            print(f'pb_vlans {s}: added VLAN {vid} ports {sorted(ports)} prio {prio}')
    wb.save(vpath)
